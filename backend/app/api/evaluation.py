from fastapi import APIRouter
from backend.app.models.evaluation import EvaluationInput
from backend.app.services.evaluation_service import process_evaluation

import csv
import io
import json
from pathlib import Path

from fastapi import File, HTTPException, UploadFile

router = APIRouter(
    prefix="/api/evaluation",
    tags=["Evaluation"]
)
SUPPORTED_REFERENCE_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".xlsx"
}

def _extract_reference_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_REFERENCE_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported format. Use PDF, DOCX, TXT, MD, CSV, JSON, or XLSX.",
        )
    try:
        if extension in {".txt", ".md"}:
            return content.decode("utf-8-sig")
        if extension == ".json":
            data = json.loads(content.decode("utf-8-sig"))
            return json.dumps(data, ensure_ascii=False, indent=2)
        if extension == ".csv":
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            return "\n".join(" | ".join(row) for row in rows)
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if extension == ".docx":
            from docx import Document
            document = Document(io.BytesIO(content))
            lines = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(lines).strip()
        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines = []
            for sheet in workbook.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(v) for v in row if v is not None]
                    if values:
                        lines.append(" | ".join(values))
            return "\n".join(lines).strip()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to extract document text: {exc}") from exc
    return ""

@router.post("/upload-reference")
async def upload_reference_document(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Please select a reference document.")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The reference document is empty.")
    text = _extract_reference_text(file.filename, content)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No readable text was found in the reference document.")
    return {"status": "success", "filename": file.filename, "text": text, "characters": len(text)}



@router.post("/submit")
def submit_evaluation(
    data: EvaluationInput
):
    return process_evaluation(data)